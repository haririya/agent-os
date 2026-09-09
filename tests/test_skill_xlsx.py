"""xlsx skill — load, eligibility, and create→inspect→edit round-trip."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

from agentos.skills.eligibility import EligibilityContext, check_eligibility
from agentos.skills.loader import SkillLoader

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"
SCRIPTS = BUNDLED / "xlsx" / "scripts"


def _spec() -> object:
    return SkillLoader(bundled_dir=BUNDLED).get_by_name("xlsx")


def test_skill_loads() -> None:
    spec = _spec()
    assert spec is not None
    assert spec.name == "xlsx"
    assert spec.metadata is not None


def test_eligibility_with_python(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: "/usr/bin/python3" if name in {"python", "python3"} else None,
    )
    spec = _spec()
    assert spec is not None
    assert check_eligibility(spec, EligibilityContext.auto())


@pytest.mark.parametrize(
    "merged",
    [
        pytest.param([{"range": "A1:B1"}], id="dictionary"),
        pytest.param(["A1:B1"], id="string"),
    ],
)
def test_round_trip_with_formula_and_merge(tmp_path: Path, merged: list[object]) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    spec = {
        "sheets": [
            {
                "name": "Sales",
                "rows": [
                    ["Region", "Revenue"],
                    ["NA", 1_200_000],
                    ["EU", 850_000],
                    ["Total", "=SUM(B2:B3)"],
                ],
                "merged": merged,
                "freeze": "A2",
            }
        ]
    }
    src = tmp_path / "book.xlsx"
    create_xlsx.build(spec).save(str(src))
    assert src.exists()

    inspected = inspect_xlsx.inspect(src, data_only=False)
    sheet = next(s for s in inspected["sheets"] if s["name"] == "Sales")
    assert sheet["max_row"] == 4
    assert sheet["max_col"] == 2
    assert sheet["merged"] == ["A1:B1"]
    assert sheet["freeze"] == "A2"

    last_row = sheet["rows"][3]
    assert last_row[1]["type"] == "f"
    assert last_row[1]["value"] == "=SUM(B2:B3)"

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(
        wb,
        [
            {"op": "set_cell", "sheet": "Sales", "row": 2, "col": 1, "value": "Americas"},
            {"op": "rename_sheet", "old": "Sales", "new": "Q3"},
        ],
    )
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    re_inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet_q3 = next(s for s in re_inspected["sheets"] if s["name"] == "Q3")
    assert sheet_q3["rows"][1][0]["value"] == "Americas"


@pytest.mark.parametrize(
    ("merged", "expected"),
    [
        pytest.param(["A1:B1", {"range": "A3:B3"}], ["A1:B1", "A3:B3"], id="mixed"),
        pytest.param([], [], id="empty"),
        pytest.param(None, [], id="null"),
        pytest.param([None, 42, {}, {"other": "A1:B1"}], [], id="unsupported-entries"),
    ],
)
def test_create_merge_formats(tmp_path: Path, merged: object, expected: list[str]) -> None:
    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build
    from agentos.skills.bundled.xlsx.scripts.inspect_xlsx import inspect

    path = tmp_path / "merges.xlsx"
    wb = build({"sheets": [{"name": "Merges", "merged": merged}, {"name": "Plain"}]})
    wb.save(path)
    wb.close()

    sheets = inspect(path, data_only=False)["sheets"]
    assert sheets[0]["merged"] == expected
    assert sheets[1]["merged"] == []


@pytest.mark.parametrize("merged", [["not-a-range"], [{"range": "not-a-range"}]])
def test_create_rejects_invalid_merge_ranges(merged: list[object]) -> None:
    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build

    with pytest.raises(ValueError, match="not a valid coordinate or range"):
        build({"sheets": [{"merged": merged}]})


def test_create_accepts_inspected_merge_ranges(tmp_path: Path) -> None:
    from openpyxl import Workbook

    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build
    from agentos.skills.bundled.xlsx.scripts.inspect_xlsx import inspect

    original = Workbook()
    ws = original.active
    assert ws is not None
    ws.merge_cells("A1:C1")
    ws.merge_cells("B3:B5")
    source = tmp_path / "source.xlsx"
    original.save(source)
    original.close()
    merges = inspect(source, data_only=False)["sheets"][0]["merged"]
    assert merges == ["A1:C1", "B3:B5"]

    # Reuse only merge metadata: inspector rows have a different schema.
    rebuilt = build({"sheets": [{"merged": merges}]})
    target = tmp_path / "rebuilt.xlsx"
    rebuilt.save(target)
    rebuilt.close()
    assert inspect(target, data_only=False)["sheets"][0]["merged"] == merges


def test_text_escapes_formula(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(
        wb,
        [
            {
                "op": "set_cell",
                "sheet": "S",
                "row": 2,
                "col": 1,
                "value": "=hello",
                "as_text": True,
            },
        ],
    )
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet = inspected["sheets"][0]
    cell_value = sheet["rows"][1][0]["value"]
    assert isinstance(cell_value, str)
    assert cell_value.lstrip("'") == "=hello"


def test_inspect_xlsx_creates_parent_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))

    out = tmp_path / "nested" / "dir" / "out.json"
    monkeypatch.setattr(sys, "argv", ["inspect_xlsx.py", str(src), "--out", str(out)])
    assert inspect_xlsx.main() == 0
    assert out.is_file()


# ── Issue #1260: set_cell with explicit null clears cell value ──────────────────


def test_set_cell_explicit_null_clears_cell(tmp_path: Path) -> None:
    """Explicit null value in set_cell operation clears cell contents upon save/reload."""
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "Sales", "rows": [["NA", 100]]}]}).save(str(src))

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    applied = edit_xlsx.apply_ops(
        wb,
        [
            {"op": "set_cell", "sheet": "Sales", "row": 1, "col": 2, "value": None},
        ],
    )
    assert applied == 1

    out = tmp_path / "out.xlsx"
    wb.save(str(out))
    wb.close()

    # Re-open from disk and inspect to verify cell is cleared
    reloaded = load_workbook(str(out))
    assert reloaded["Sales"].cell(row=1, column=2).value is None
    reloaded.close()

    inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet = inspected["sheets"][0]
    row = sheet["rows"][0]
    # Row list length should be 1 since cell (1, 2) is cleared (empty)
    assert len(row) == 1 or row[1]["value"] is None


def test_set_cell_missing_value_key_skipped(tmp_path: Path) -> None:
    """Operation with missing 'value' key is skipped and leaves existing cell unchanged."""
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "Sales", "rows": [["NA", 100]]}]}).save(str(src))

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    applied = edit_xlsx.apply_ops(
        wb,
        [
            {"op": "set_cell", "sheet": "Sales", "row": 1, "col": 2},  # missing 'value' key
        ],
    )
    assert applied == 0

    out = tmp_path / "out.xlsx"
    wb.save(str(out))
    wb.close()

    reloaded = load_workbook(str(out))
    assert reloaded["Sales"].cell(row=1, column=2).value == 100
    reloaded.close()


def test_set_cell_explicit_null_via_cli(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Round-trip via edit_xlsx.py CLI script with JSON null value clears cell."""
    import json

    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "Sales", "rows": [["NA", 100]]}]}).save(str(src))

    ops_file = tmp_path / "ops.json"
    ops_file.write_text(
        json.dumps([{"op": "set_cell", "sheet": "Sales", "row": 1, "col": 2, "value": None}]),
        encoding="utf-8",
    )
    out_file = tmp_path / "out.xlsx"

    monkeypatch.setattr(
        sys, "argv", ["edit_xlsx.py", str(src), str(ops_file), "--out", str(out_file)]
    )
    assert edit_xlsx.main() == 0

    from openpyxl import load_workbook

    reloaded = load_workbook(str(out_file))
    assert reloaded["Sales"].cell(row=1, column=2).value is None
    reloaded.close()
